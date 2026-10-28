"""An organization's hiring data as Excel, CSV (zipped, one file per sheet)
or a PDF summary, so a trial can keep a copy before it ends or is purged.

Everything is read with organization_id from the signed-in admin, never
from the request. Resume text, AI vectors and internal IDs aren't exported.
"""
import csv
import html
import io
import zipfile
from datetime import datetime, timezone

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.application import Application
from app.models.candidate import Candidate
from app.models.interview import Interview
from app.models.job import Job
from app.models.offer import Offer
from app.models.organization import Organization
from app.models.scorecard import Scorecard

FORMATS = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "application/zip",
    "pdf": "application/pdf",
}


def _d(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    return str(value)


def _cell(value) -> str:
    """Spreadsheet apps run a cell starting with = + - @ as a formula."""
    text = _d(value)
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


async def collect(db: AsyncSession, organization_id) -> dict[str, list[dict]]:
    org = organization_id
    jobs = (await db.execute(
        select(Job).where(Job.organization_id == org, Job.is_deleted.is_(False)).order_by(Job.created_at)
    )).scalars().all()
    job_titles = {j.id: j.title for j in jobs}
    candidates = (await db.execute(
        select(Candidate).where(Candidate.organization_id == org, Candidate.is_deleted.is_(False)).order_by(Candidate.created_at)
    )).scalars().all()
    names = {c.id: c.full_name for c in candidates}
    apps = (await db.execute(
        select(Application).where(Application.organization_id == org).order_by(Application.applied_at)
    )).scalars().all()
    app_ids = {a.id: a for a in apps}
    interviews = (await db.execute(
        select(Interview).where(Interview.organization_id == org).order_by(Interview.scheduled_at)
    )).scalars().all()
    scorecards = (await db.execute(
        select(Scorecard).where(Scorecard.organization_id == org).order_by(Scorecard.submitted_at)
    )).scalars().all()
    offers = (await db.execute(
        select(Offer).where(Offer.organization_id == org).order_by(Offer.created_at)
    )).scalars().all()
    interview_titles = {i.id: i.title for i in interviews}

    def job_of(application_id):
        a = app_ids.get(application_id)
        return job_titles.get(a.job_id, "") if a else ""

    def cand_of(application_id):
        a = app_ids.get(application_id)
        return names.get(a.candidate_id, "") if a else ""

    return {
        "Jobs": [{
            "Title": j.title, "Status": j.status, "Location": j.location, "Type": j.job_type,
            "Experience level": j.experience_level, "Remote": "Yes" if j.is_remote else "No",
            "Openings": j.openings, "Skills": j.skills_required, "Created": j.created_at,
        } for j in jobs],
        "Candidates": [{
            "Name": c.full_name, "Email": c.email, "Phone": c.phone, "Location": c.location,
            "Current title": c.current_title, "Current company": c.current_company,
            "Experience (years)": c.years_experience, "Skills": c.skills, "LinkedIn": c.linkedin_url,
            "Notice period (days)": c.notice_period_days, "Tags": c.tags, "Added": c.created_at,
        } for c in candidates],
        "Pipeline": [{
            "Candidate": names.get(a.candidate_id, ""), "Job": job_titles.get(a.job_id, ""),
            "Stage": (a.stage or "").replace("_", " ").title(),
            "Match score": f"{a.match_score:.1f}%" if a.match_score is not None else "",
            "Source": a.source, "Applied": a.applied_at, "Stage changed": a.stage_changed_at,
            "Rejection reason": a.rejection_reason,
        } for a in apps],
        "Interviews": [{
            "Title": i.title, "Candidate": names.get(i.candidate_id, ""), "Job": job_of(i.application_id),
            "Type": i.interview_type, "Status": i.status, "Scheduled": i.scheduled_at,
            "Duration (min)": i.duration_minutes, "Feedback": i.feedback,
        } for i in interviews],
        "Scorecards": [{
            "Interview": interview_titles.get(s.interview_id, ""), "Candidate": cand_of(s.application_id),
            "Rating": s.overall_rating, "Recommendation": s.recommendation, "Strengths": s.strengths,
            "Weaknesses": s.weaknesses, "Summary": s.summary, "Submitted": s.submitted_at,
        } for s in scorecards],
        "Offers": [{
            "Candidate": cand_of(o.application_id), "Job": job_of(o.application_id), "Position": o.position_title,
            "Status": o.status, "Base salary": o.base_salary, "Currency": o.salary_currency,
            "Start date": o.start_date, "Sent": o.sent_at, "Responded": o.responded_at,
        } for o in offers],
    }


def to_xlsx(data: dict[str, list[dict]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for sheet, rows in data.items():
        ws = wb.create_sheet(sheet)
        if not rows:
            ws.append(["No records"])
            continue
        headers = list(rows[0])
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([_cell(row[h]) for h in headers])
        for i, h in enumerate(headers, 1):
            width = max(len(h), *(len(_d(r[h])) for r in rows))
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(60, width + 2)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_csv_zip(data: dict[str, list[dict]]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for sheet, rows in data.items():
            buf = io.StringIO()
            writer = csv.writer(buf)
            if rows:
                headers = list(rows[0])
                writer.writerow(headers)
                writer.writerows([_cell(r[h]) for h in headers] for r in rows)
            # Excel opens UTF-8 CSVs correctly only with a BOM.
            zf.writestr(f"{sheet.lower()}.csv", "﻿" + buf.getvalue())
    return out.getvalue()


def to_pdf(data: dict[str, list[dict]], org_name: str) -> bytes:
    from xhtml2pdf import pisa

    e = html.escape
    counts = "".join(f"<tr><td>{e(k)}</td><td>{len(v)}</td></tr>" for k, v in data.items())
    stage_counts: dict[str, dict[str, int]] = {}
    for row in data["Pipeline"]:
        stage_counts.setdefault(row["Job"] or "Unassigned", {}).setdefault(row["Stage"] or "-", 0)
        stage_counts[row["Job"] or "Unassigned"][row["Stage"] or "-"] += 1
    pipeline = "".join(
        f"<h3>{e(job)}</h3><table>" + "".join(f"<tr><td>{e(s)}</td><td>{n}</td></tr>" for s, n in stages.items()) + "</table>"
        for job, stages in stage_counts.items()
    ) or "<p>No candidates in a pipeline yet.</p>"
    jobs = "".join(
        f"<tr><td>{e(_d(j['Title']))}</td><td>{e(_d(j['Status']))}</td><td>{e(_d(j['Location']))}</td><td>{e(_d(j['Created']))}</td></tr>"
        for j in data["Jobs"]
    )
    candidates = "".join(
        f"<tr><td>{e(_d(c['Name']))}</td><td>{e(_d(c['Email']))}</td><td>{e(_d(c['Current title']))}</td><td>{e(_d(c['Added']))}</td></tr>"
        for c in data["Candidates"]
    )
    doc = f"""<html><head><style>
        @page {{ size: A4; margin: 1.6cm; }}
        body {{ font-family: Helvetica; font-size: 9pt; color: #202124; }}
        h1 {{ font-size: 16pt; margin: 0; }} h2 {{ font-size: 12pt; margin-top: 18pt; border-bottom: 1px solid #dadce0; }}
        h3 {{ font-size: 10pt; margin: 8pt 0 2pt; }}
        table {{ width: 100%; border-collapse: collapse; }} td, th {{ border-bottom: 1px solid #eee; padding: 3pt; text-align: left; }}
        th {{ background: #f1f3f4; }} .muted {{ color: #70757a; }}
    </style></head><body>
        <h1>{e(org_name)}: hiring data export</h1>
        <p class="muted">Generated {_d(datetime.now(timezone.utc))} from Hybent Hiring. The Excel and CSV exports hold every field.</p>
        <h2>Summary</h2><table><tr><th>Records</th><th>Count</th></tr>{counts}</table>
        <h2>Jobs</h2><table><tr><th>Title</th><th>Status</th><th>Location</th><th>Created</th></tr>{jobs}</table>
        <h2>Pipeline by job</h2>{pipeline}
        <h2>Candidates</h2><table><tr><th>Name</th><th>Email</th><th>Current title</th><th>Added</th></tr>{candidates}</table>
    </body></html>"""
    out = io.BytesIO()
    result = pisa.CreatePDF(doc, dest=out)
    if result.err:
        raise RuntimeError("PDF generation failed")
    return out.getvalue()


async def build(db: AsyncSession, organization_id, fmt: str) -> tuple[bytes, str, str]:
    """(file bytes, media type, filename)."""
    data = await collect(db, organization_id)
    org = await db.get(Organization, organization_id)
    org_name = org.name if org else "Organization"
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d")
    slug = (org.slug if org else "hybent")[:40]
    if fmt == "xlsx":
        body = to_xlsx(data)
    elif fmt == "csv":
        body = to_csv_zip(data)
    else:
        body = to_pdf(data, org_name)
    ext = "zip" if fmt == "csv" else fmt
    return body, FORMATS[fmt], f"hybent-export-{slug}-{stamp}.{ext}"
