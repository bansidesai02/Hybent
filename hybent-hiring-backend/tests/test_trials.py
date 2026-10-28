"""The 7-day free trial: request → super admin approval → reminders → export
→ expiry → paid upgrade or purge. Stripe, email, DNS and Celery are faked."""
import io
import types
import uuid
import zipfile
from datetime import datetime, timedelta, timezone

import pytest
import pytest_asyncio
from sqlalchemy import func, select

import app.services.payment_service as payment_service
import app.services.trial_service as trial_service
from app.models.candidate import Candidate
from app.models.job import Job
from app.models.organization import Organization
from app.models.super_admin import CompanyFeatureFlag, CompanySubscription, SubscriptionPlan, SuperAdminAuditLog
from app.models.trial_request import TrialRequest
from app.models.user import User
from app.services import email_service, lead_checks
from conftest import _make_user, auth_headers
from test_payments import FakeStripe, fake_stripe  # noqa: F401  fake_stripe is a fixture

CODES: dict[str, str] = {}
# The real send_email, before the outbox fixture replaces it.
REAL_SEND_EMAIL = email_service.send_email

TRIAL_FEATURES = {"ai": True, "video": True, "bulk": True, "domain": False, "analytics": True, "copilot": True}

VALID = {
    "company_name": "Initech Systems",
    "company_website": "https://www.initech-example.com",
    "company_linkedin_url": "https://www.linkedin.com/company/initech",
    "company_size": "51-200",
    "industry": "Software / SaaS",
    "country": "United States",
    "hires_per_year": "11-50",
    "contact_name": "Bill Lumbergh",
    "contact_title": "VP Engineering",
    "contact_email": "bill@initech-example.com",
    "contact_phone": "+1 415 555 2671",
    "contact_linkedin_url": "https://www.linkedin.com/in/bill-lumbergh",
    "hiring_needs": "Five backend engineers this quarter <script>",
    "accept_terms": True,
    "elapsed_ms": 9000,
}


@pytest.fixture(autouse=True)
def outbox(monkeypatch):
    """Captured emails and notifications; no DNS, rate limit or broker."""
    sent = {"emails": [], "notifications": []}
    monkeypatch.setattr(email_service, "send_email", lambda to, subject, body, **k: sent["emails"].append((to, subject, body)) or True)
    monkeypatch.setattr(payment_service, "send_email", lambda to, subject, body: sent["emails"].append((to, subject, body)) or True)

    async def accepts(domain):
        return True

    monkeypatch.setattr(lead_checks, "domain_accepts_mail", accepts)
    monkeypatch.setattr(lead_checks, "rate_limited", lambda ip: False)

    async def no_storage(org_id, candidate_id):
        sent.setdefault("deleted_files", []).append((org_id, candidate_id))

    import app.services.supabase_storage_service as storage
    monkeypatch.setattr(storage, "delete_candidate_files", no_storage)

    # OTP codes are captured, not emailed.
    def capture_code(to, code, minutes):
        CODES[to] = code
        return True

    monkeypatch.setattr(email_service, "send_trial_otp_email", capture_code)
    # No AI in tests: drafts use the fixed templates.
    monkeypatch.setattr(trial_service.settings, "groq_api_key", "")

    fake = types.SimpleNamespace(delay=lambda *a, **k: sent["notifications"].append(a))
    import app.services.ai_credit_service as credit_svc
    import app.tasks.notifications as notifications
    for mod in (credit_svc, notifications, payment_service):
        monkeypatch.setattr(mod, "notify_organization_roles", fake)
    return sent


@pytest_asyncio.fixture
async def plans(db_session):
    rows = [
        SubscriptionPlan(name="Standard", term_months=1, price_monthly=29.0, price_yearly=348.0, sort_order=1,
                         ai_credits_monthly=5000, included_recruiters=0, features={**TRIAL_FEATURES, "copilot": False}),
        SubscriptionPlan(name="Pro", term_months=1, price_monthly=39.0, price_yearly=468.0, sort_order=2,
                         ai_credits_monthly=10000, included_recruiters=1, features=TRIAL_FEATURES),
        SubscriptionPlan(name="Trial Standard", term_months=1, price_monthly=0.0, price_yearly=0.0, sort_order=98,
                         ai_credits_monthly=500, included_admins=1, included_recruiters=0,
                         features={**TRIAL_FEATURES, "copilot": False}),
        SubscriptionPlan(name="Trial Pro", term_months=1, price_monthly=0.0, price_yearly=0.0, sort_order=99,
                         ai_credits_monthly=500, included_admins=1, included_recruiters=1, features=TRIAL_FEATURES),
    ]
    db_session.add_all(rows)
    await db_session.commit()
    return {p.name: p for p in rows}


async def _verify(client, email):
    """The applicant gets a code and enters it; returns the response or the token."""
    res = await client.post("/api/public/trial-request/send-code", json={"email": email})
    if res.status_code != 200:
        return res
    res = await client.post("/api/public/trial-request/verify-code", json={"email": email, "code": CODES[email.lower()]})
    assert res.status_code == 200, res.text
    return res.json()["data"]["verification_token"]


async def _submit(client, **change):
    payload = {**VALID, **change}
    token = await _verify(client, payload["contact_email"])
    if not isinstance(token, str):
        return token
    return await client.post("/api/public/trial-request", json={**payload, "email_verification_token": token})


async def _send(client, req_id, headers, kind="approve", **extra):
    return await client.post(f"/v1/super-admin/trials/{req_id}/send", headers=headers, json={
        "kind": kind, "subject": "Your Hybent Hiring trial", "body": "Hi Bill,\n\nThanks for your request.", **extra,
    })


async def _approved(client, db_session, super_admin_headers, **change):
    """Submit and approve a request; returns (request, org admin user)."""
    res = await _submit(client, **change)
    assert res.status_code == 200, res.text
    email = change.get("contact_email", VALID["contact_email"])
    req = (await db_session.execute(select(TrialRequest).where(TrialRequest.contact_email == email))).scalar_one()
    res = await _send(client, req.id, super_admin_headers)
    assert res.status_code == 200, res.text
    await db_session.refresh(req)
    admin = (await db_session.execute(select(User).where(User.email == email))).scalar_one()
    return req, admin


async def _set_trial_end(db_session, req: TrialRequest, when: datetime):
    sub = (await db_session.execute(
        select(CompanySubscription).where(CompanySubscription.organization_id == req.organization_id)
    )).scalar_one()
    req.trial_end = when
    sub.trial_end = when
    sub.current_period_end = when
    await db_session.commit()


# ── Request ──────────────────────────────────────────────────────────────────

async def test_valid_request_is_saved_and_emailed(client, db_session, outbox):
    res = await _submit(client)
    assert res.status_code == 200, res.text
    req = (await db_session.execute(select(TrialRequest))).scalar_one()
    assert req.status == "pending"
    assert req.company_domain == "initech-example.com"
    assert req.contact_phone.startswith("+1")
    assert req.risk_flags == []
    assert req.email_verified_at is not None
    # Only the team is emailed; the applicant got their code and sees the thank-you.
    recipients = [to for to, _, _ in outbox["emails"]]
    assert recipients == ["info@hybent.com"]
    alert = outbox["emails"][0][2]
    assert "<script>" not in alert and "&lt;script&gt;" in alert


@pytest.mark.parametrize("email", ["bill@gmail.com", "bill@yahoo.com", "bill@outlook.com", "bill@hotmail.com", "bill@mailinator.com"])
async def test_personal_and_disposable_emails_are_refused(client, db_session, email):
    res = await _submit(client, contact_email=email)
    assert res.status_code == 422
    assert res.json()["details"]["field"] == "contact_email"
    assert await db_session.scalar(select(func.count(TrialRequest.id))) == 0


@pytest.mark.parametrize("change", [
    {"company_linkedin_url": "https://www.linkedin.com/in/someone"},
    {"contact_linkedin_url": "https://www.linkedin.com/company/initech"},
    {"company_website": "not a site"},
    {"company_size": "huge"},
    {"industry": "Basket weaving"},
    {"accept_terms": False},
    {"company_name": ""},
])
async def test_invalid_fields_are_refused(client, db_session, change):
    assert (await _submit(client, **change)).status_code == 422
    assert await db_session.scalar(select(func.count(TrialRequest.id))) == 0


async def test_mismatched_domain_is_flagged_not_refused(client, db_session):
    res = await _submit(client, company_website="https://initech.io")
    assert res.status_code == 200
    req = (await db_session.execute(select(TrialRequest))).scalar_one()
    assert req.risk_flags and "doesn't match" in req.risk_flags[0]


async def test_bots_are_dropped_quietly(client, db_session, outbox):
    assert (await _submit(client, homepage="http://spam")).status_code == 200
    assert (await _submit(client, elapsed_ms=200, contact_email="fast@initech-example.com")).status_code == 200
    assert await db_session.scalar(select(func.count(TrialRequest.id))) == 0
    assert outbox["emails"] == []  # no alert to the team


async def test_one_open_request_per_domain(client):
    assert (await _submit(client)).status_code == 200
    assert (await _submit(client)).status_code == 409
    # A colleague at the same company can't open a second one.
    res = await _submit(client, contact_email="peter@initech-example.com")
    assert res.status_code == 409


async def test_existing_account_email_is_refused(client, admin_user):
    domain = admin_user.email.split("@")[1]
    assert domain == "example.com"
    res = await _submit(client, contact_email=admin_user.email, company_website="https://example.com")
    assert res.status_code == 409


# ── Review ───────────────────────────────────────────────────────────────────

async def test_only_super_admins_review(client, admin_headers, recruiter_headers):
    for headers in (admin_headers, recruiter_headers, {}):
        res = await client.get("/v1/super-admin/trials", headers=headers)
        assert res.status_code in (401, 403)
    res = await _send(client, uuid.uuid4(), admin_headers)
    assert res.status_code == 403
    res = await client.post(f"/v1/super-admin/trials/{uuid.uuid4()}/draft", headers=admin_headers, json={"kind": "approve"})
    assert res.status_code == 403


async def test_approval_creates_a_trial_organization(client, db_session, plans, super_admin_headers, outbox):
    before = datetime.now(timezone.utc)
    req, admin = await _approved(client, db_session, super_admin_headers)

    assert req.status == "active"
    assert timedelta(days=7) - timedelta(minutes=1) < req.trial_end - before < timedelta(days=7, minutes=1)
    org = await db_session.get(Organization, req.organization_id)
    assert org.is_active and org.name == "Initech Systems"
    assert admin.role == "admin" and admin.is_active and admin.organization_id == org.id
    sub = (await db_session.execute(select(CompanySubscription).where(CompanySubscription.organization_id == org.id))).scalar_one()
    assert sub.is_trial and sub.plan_id == plans["Trial Pro"].id and sub.trial_end == req.trial_end
    flags = dict((await db_session.execute(
        select(CompanyFeatureFlag.flag_key, CompanyFeatureFlag.is_enabled).where(CompanyFeatureFlag.organization_id == org.id)
    )).all())
    assert flags == {**TRIAL_FEATURES, "ai_calls": False}
    audit = (await db_session.execute(select(SuperAdminAuditLog).where(SuperAdminAuditLog.module == "Trials"))).scalar_one()
    assert "Approved" in audit.action

    to, subject, body = outbox["emails"][-1]
    assert to == "bill@initech-example.com" and subject == "Your Hybent Hiring trial"
    assert f"/setup?token={req.setup_token}" in body and "Your trial ends on" in body
    assert "Set up your organization" in body

    # Approving twice does nothing more.
    assert (await _send(client, req.id, super_admin_headers)).status_code == 409
    # The company can't request another trial.
    res = await _submit(client, contact_email="peter@initech-example.com")
    assert res.status_code == 409


async def test_listing_counts_and_filters(client, db_session, plans, super_admin_headers):
    await _submit(client)
    await _submit(client, contact_email="ann@globex-example.com", company_website="https://globex-example.com",
                  company_name="Globex")
    res = await client.get("/v1/super-admin/trials?status=pending&search=globex", headers=super_admin_headers)
    data = res.json()["data"]
    assert data["total"] == 1 and data["items"][0]["company_name"] == "Globex"
    assert data["counts"]["pending"] == 2


async def test_rejection_emails_the_applicant(client, db_session, plans, super_admin_headers, outbox):
    await _submit(client)
    req = (await db_session.execute(select(TrialRequest))).scalar_one()
    res = await _send(client, req.id, super_admin_headers, kind="reject",
                      body="Hi Bill,\n\nWe couldn't verify the company.", reason="Unverified company")
    assert res.status_code == 200, res.text
    await db_session.refresh(req)
    assert req.status == "rejected" and req.organization_id is None and req.rejection_reason == "Unverified company"
    assert outbox["emails"][-1][0] == "bill@initech-example.com"
    assert "verify the company" in outbox["emails"][-1][2] and "/setup" not in outbox["emails"][-1][2]
    assert (await _send(client, req.id, super_admin_headers)).status_code == 409
    detail = (await client.get(f"/v1/super-admin/trials/{req.id}", headers=super_admin_headers)).json()["data"]
    assert [m["kind"] for m in detail["messages"]] == ["reject"] and detail["messages"][0]["delivered"]


# ── Access ───────────────────────────────────────────────────────────────────

async def test_expired_trial_limits_access_exactly_at_trial_end(client, db_session, plans, super_admin_headers):
    req, admin = await _approved(client, db_session, super_admin_headers)
    org = await db_session.get(Organization, req.organization_id)
    recruiter = await _make_user(db_session, org, "recruiter")
    h_admin, h_rec = auth_headers(admin), auth_headers(recruiter)

    assert (await client.get("/v1/jobs", headers=h_admin)).status_code == 200
    assert (await client.get("/v1/jobs", headers=h_rec)).status_code == 200
    status = (await client.get("/v1/trial/status", headers=h_admin)).json()["data"]
    assert status["status"] == "active" and status["seconds_left"] > 6 * 86400

    # No job has run: access still ends the moment trial_end passes.
    await _set_trial_end(db_session, req, datetime.now(timezone.utc) - timedelta(seconds=1))
    res = await client.get("/v1/jobs", headers=h_admin)
    assert res.status_code == 402 and res.json()["details"]["code"] == "trial_expired"
    assert (await client.get("/v1/billing", headers=h_admin)).status_code == 200
    assert (await client.get("/v1/trial/status", headers=h_admin)).json()["data"]["status"] == "expired"
    assert (await client.get("/v1/notifications", headers=h_admin)).status_code == 200
    res = await client.get("/v1/billing", headers=h_rec)
    assert res.status_code == 403 and res.json()["details"]["code"] == "trial_expired"


async def test_old_trial_end_on_regular_clients_is_not_enforced(client, db_session, plans, organization, admin_headers):
    """trial_end on subscriptions made before this feature never cut access, and still doesn't."""
    past = datetime.now(timezone.utc) - timedelta(days=30)
    db_session.add(CompanySubscription(
        organization_id=organization.id, plan_id=plans["Standard"].id, status="active",
        current_period_end=datetime.now(timezone.utc) + timedelta(days=10), trial_end=past,
    ))
    await db_session.commit()
    assert (await client.get("/v1/jobs", headers=admin_headers)).status_code == 200


async def test_members_of_expired_trial_cannot_sign_in(client, db_session, plans, super_admin_headers):
    req, admin = await _approved(client, db_session, super_admin_headers)
    org = await db_session.get(Organization, req.organization_id)
    recruiter = await _make_user(db_session, org, "recruiter")
    await _set_trial_end(db_session, req, datetime.now(timezone.utc) - timedelta(minutes=5))
    res = await client.post("/v1/auth/login", json={"email": recruiter.email, "password": "password123"})
    assert res.status_code == 403 and "trial has ended" in res.json()["message"]


# ── Lifecycle ────────────────────────────────────────────────────────────────

async def test_reminders_expiry_and_purge_run_once_each(client, db_session, plans, super_admin_headers, outbox, other_organization):
    req, admin = await _approved(client, db_session, super_admin_headers)
    org_id = req.organization_id
    cand = Candidate(organization_id=org_id, email="c1@x-example.com", full_name="Cand One")
    db_session.add(cand)
    db_session.add(Candidate(organization_id=other_organization.id, email="c2@x-example.com", full_name="Other Org Cand"))
    await db_session.commit()
    cand_id = cand.id
    outbox["emails"].clear()

    async def run(at):
        return await trial_service.process_lifecycle(db_session, now=at)

    end = req.trial_end
    assert (await run(end - timedelta(days=2)))["reminder_1d"] == 0
    assert (await run(end - timedelta(hours=23)))["reminder_1d"] == 1
    assert (await run(end - timedelta(hours=22)))["reminder_1d"] == 0
    assert "ends in 1 day" in outbox["emails"][-1][1]
    assert (await run(end - timedelta(minutes=50)))["reminder_1h"] == 1
    assert (await run(end - timedelta(minutes=40)))["reminder_1h"] == 0
    assert "ends in 1 hour" in outbox["emails"][-1][1]

    done = await run(end + timedelta(minutes=1))
    assert done["expired"] == 1
    await db_session.refresh(req)
    assert req.status == "expired" and req.data_purge_at == end + timedelta(days=30)
    assert "has ended" in outbox["emails"][-1][1]
    assert (await run(end + timedelta(minutes=6)))["expired"] == 0
    # In-app notices: 1 day, 1 hour, expired.
    assert len(outbox["notifications"]) == 3

    assert (await run(end + timedelta(days=24)))["purge_warning"] == 1
    assert (await run(end + timedelta(days=25)))["purge_warning"] == 0
    assert (await run(end + timedelta(days=29)))["purged"] == 0

    done = await run(end + timedelta(days=30, minutes=1))
    assert done["purged"] == 1
    req_id, other_id = req.id, other_organization.id
    db_session.expire_all()
    req = await db_session.get(TrialRequest, req_id)
    assert req.status == "purged" and req.organization_id is None
    assert await db_session.get(Organization, org_id) is None
    assert await db_session.scalar(select(func.count(User.id)).where(User.organization_id == org_id)) == 0
    assert await db_session.scalar(select(func.count(Candidate.id)).where(Candidate.organization_id == org_id)) == 0
    assert outbox["deleted_files"] == [(str(org_id), str(cand_id))]
    # Other tenants are untouched.
    assert await db_session.scalar(select(func.count(Candidate.id)).where(Candidate.organization_id == other_id)) == 1


async def test_extend_moves_trial_end_and_resets_reminders(client, db_session, plans, super_admin_headers):
    req, _ = await _approved(client, db_session, super_admin_headers)
    old_end = req.trial_end
    await trial_service.process_lifecycle(db_session, now=old_end - timedelta(hours=20))
    await db_session.refresh(req)
    assert req.reminder_1d_sent_at is not None
    res = await client.post(f"/v1/super-admin/trials/{req.id}/extend", headers=super_admin_headers, json={"days": 3})
    assert res.status_code == 200, res.text
    await db_session.refresh(req)
    assert req.trial_end >= old_end + timedelta(days=3) - timedelta(seconds=5)
    assert req.reminder_1d_sent_at is None


# ── Export ───────────────────────────────────────────────────────────────────

async def test_export_has_only_the_admins_organization(client, db_session, plans, super_admin_headers, other_organization):
    req, admin = await _approved(client, db_session, super_admin_headers)
    db_session.add(Job(organization_id=req.organization_id, created_by_id=admin.id, title="Backend Engineer",
                       job_type="full_time", description="Build APIs"))
    db_session.add(Candidate(organization_id=req.organization_id, email="ana@x-example.com", full_name="=HYPERLINK(evil)"))
    db_session.add(Candidate(organization_id=other_organization.id, email="leak@x-example.com", full_name="Should Not Leak"))
    await db_session.commit()
    h = auth_headers(admin)

    res = await client.get("/v1/trial/export?format=csv", headers=h)
    assert res.status_code == 200 and res.headers["content-type"] == "application/zip"
    zf = zipfile.ZipFile(io.BytesIO(res.content))
    candidates = zf.read("candidates.csv").decode("utf-8-sig")
    assert "ana@x-example.com" in candidates and "leak@" not in candidates
    assert "'=HYPERLINK" in candidates  # formula neutralised
    assert "Backend Engineer" in zf.read("jobs.csv").decode("utf-8-sig")

    res = await client.get("/v1/trial/export?format=xlsx", headers=h)
    assert res.status_code == 200 and res.content[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Jobs", "Candidates", "Pipeline", "Interviews", "Scorecards", "Offers"]

    res = await client.get("/v1/trial/export?format=pdf", headers=h)
    assert res.status_code == 200 and res.content[:4] == b"%PDF"

    # Still works after expiry; recruiters can't export.
    await _set_trial_end(db_session, req, datetime.now(timezone.utc) - timedelta(minutes=1))
    assert (await client.get("/v1/trial/export?format=csv", headers=h)).status_code == 200
    org = await db_session.get(Organization, req.organization_id)
    recruiter = await _make_user(db_session, org, "recruiter")
    assert (await client.get("/v1/trial/export", headers=auth_headers(recruiter))).status_code == 403


# ── Upgrade ──────────────────────────────────────────────────────────────────

async def test_expired_trial_upgrades_to_pro_and_keeps_its_data(client, db_session, plans, super_admin_headers, fake_stripe):
    req, admin = await _approved(client, db_session, super_admin_headers)
    db_session.add(Candidate(organization_id=req.organization_id, email="keep@x-example.com", full_name="Keep Me"))
    await db_session.commit()
    await _set_trial_end(db_session, req, datetime.now(timezone.utc) - timedelta(minutes=1))
    await trial_service.process_lifecycle(db_session)
    h = auth_headers(admin)
    assert (await client.get("/v1/jobs", headers=h)).status_code == 402

    # Only Standard or Pro.
    assert (await client.post("/v1/billing/upgrade", headers=h, json={"plan_name": "Trial Pro"})).status_code == 422
    res = await client.post("/v1/billing/upgrade", headers=h, json={"plan_name": "Pro"})
    assert res.status_code == 200, res.text
    url = res.json()["data"]["checkout_url"]
    sid = url.rsplit("/", 1)[1]
    session = fake_stripe.complete(sid, amount_cents=3900)
    token = session["success_url"].split("/pay/")[1].split("?")[0]
    res = await client.post(f"/v1/payments/{token}/confirm", json={"session_id": sid})
    assert res.status_code == 200, res.text

    req_id, org_id, pro_id = req.id, req.organization_id, plans["Pro"].id
    db_session.expire_all()
    sub = (await db_session.execute(
        select(CompanySubscription).where(CompanySubscription.organization_id == org_id)
    )).scalar_one()
    assert not sub.is_trial and sub.status == "active" and sub.plan_id == pro_id and sub.trial_end is None
    req = await db_session.get(TrialRequest, req_id)
    assert req.status == "converted" and req.data_purge_at is None
    assert (await client.get("/v1/jobs", headers=h)).status_code == 200
    # Nothing is purged later.
    done = await trial_service.process_lifecycle(db_session, now=datetime.now(timezone.utc) + timedelta(days=40))
    assert done["purged"] == 0
    assert await db_session.scalar(select(func.count(Candidate.id)).where(Candidate.organization_id == org_id)) == 1


async def test_paid_clients_cannot_use_trial_upgrade(client, db_session, plans, organization, admin_headers, fake_stripe):
    db_session.add(CompanySubscription(
        organization_id=organization.id, plan_id=plans["Standard"].id, status="active",
        current_period_end=datetime.now(timezone.utc) + timedelta(days=10),
    ))
    await db_session.commit()
    res = await client.post("/v1/billing/upgrade", headers=admin_headers, json={"plan_name": "Pro"})
    assert res.status_code == 409


async def test_trial_plan_is_not_listed_for_sale(client, db_session, plans, organization, admin_headers):
    db_session.add(CompanySubscription(
        organization_id=organization.id, plan_id=plans["Standard"].id, status="active",
        current_period_end=datetime.now(timezone.utc) + timedelta(days=10),
    ))
    await db_session.commit()
    names = [p["name"] for p in (await client.get("/v1/billing", headers=admin_headers)).json()["data"]["plans"]]
    assert not any(n.startswith("Trial") for n in names) and names == ["Standard", "Pro"]


async def test_standard_trial_gets_standard_features(client, db_session, plans, super_admin_headers, outbox):
    req, _ = await _approved(client, db_session, super_admin_headers, plan="standard")
    assert req.plan_interest == "standard"
    sub = (await db_session.execute(
        select(CompanySubscription).where(CompanySubscription.organization_id == req.organization_id)
    )).scalar_one()
    assert sub.is_trial and sub.plan_id == plans["Trial Standard"].id
    flags = dict((await db_session.execute(
        select(CompanyFeatureFlag.flag_key, CompanyFeatureFlag.is_enabled)
        .where(CompanyFeatureFlag.organization_id == req.organization_id)
    )).all())
    assert flags["copilot"] is False
    assert "/setup?token=" in outbox["emails"][-1][2]


# ── Email verification ───────────────────────────────────────────────────────

async def test_request_needs_a_verified_email(client, db_session):
    res = await client.post("/api/public/trial-request", json=VALID)
    assert res.status_code == 422 and res.json()["details"]["field"] == "contact_email"
    # A token for another address doesn't count.
    token = await _verify(client, "other@initech-example.com")
    res = await client.post("/api/public/trial-request", json={**VALID, "email_verification_token": token})
    assert res.status_code == 422
    assert await db_session.scalar(select(func.count(TrialRequest.id))) == 0


async def test_wrong_codes_are_limited_and_tokens_single_use(client, db_session):
    email = VALID["contact_email"]
    assert (await client.post("/api/public/trial-request/send-code", json={"email": email})).status_code == 200
    # Asking again straight away is refused.
    assert (await client.post("/api/public/trial-request/send-code", json={"email": email})).status_code == 429
    for _ in range(4):
        res = await client.post("/api/public/trial-request/verify-code", json={"email": email, "code": "000000"})
        assert res.status_code == 422 and res.json()["details"]["field"] == "otp"
    res = await client.post("/api/public/trial-request/verify-code", json={"email": email, "code": "000000"})
    assert "Too many" in res.json()["message"]
    # Even the right code fails now.
    res = await client.post("/api/public/trial-request/verify-code", json={"email": email, "code": CODES[email]})
    assert res.status_code == 422


async def test_personal_email_gets_no_code(client):
    res = await client.post("/api/public/trial-request/send-code", json={"email": "bill@gmail.com"})
    assert res.status_code == 422 and "bill@gmail.com" not in CODES


# ── Replies ──────────────────────────────────────────────────────────────────

async def test_draft_falls_back_to_template_without_ai(client, db_session, plans, super_admin_headers):
    await _submit(client)
    req = (await db_session.execute(select(TrialRequest))).scalar_one()
    res = await client.post(f"/v1/super-admin/trials/{req.id}/draft", headers=super_admin_headers, json={"kind": "approve"})
    data = res.json()["data"]
    assert res.status_code == 200 and data["ai_generated"] is False
    assert "Bill" in data["body"] and "Initech Systems" in data["subject"]


async def test_info_email_keeps_request_pending(client, db_session, plans, super_admin_headers, outbox):
    await _submit(client)
    req = (await db_session.execute(select(TrialRequest))).scalar_one()
    res = await _send(client, req.id, super_admin_headers, kind="info", body="Hi Bill,\n\nCould you share <your> company registration?")
    assert res.status_code == 200
    await db_session.refresh(req)
    assert req.status == "pending"
    assert "&lt;your&gt;" in outbox["emails"][-1][2]


# ── Organization setup ───────────────────────────────────────────────────────

SETUP = {
    "password": "Initech2026", "full_name": "Bill Lumbergh", "company_name": "Initech",
    "website": "https://initech-example.com", "industry": "Software / SaaS", "size": "51-200",
    "timezone": "America/New_York", "brand_color": "#ff5500", "email_sender_name": "Initech Talent",
}


async def test_setup_link_sets_up_the_organization_once(client, db_session, plans, super_admin_headers):
    req, admin = await _approved(client, db_session, super_admin_headers)
    token = req.setup_token
    info = (await client.get(f"/api/public/trial-setup/{token}")).json()["data"]
    assert info["email"] == "bill@initech-example.com" and info["plan"] == "Pro"

    assert (await client.post(f"/api/public/trial-setup/{token}", json={**SETUP, "password": "short"})).status_code == 422
    res = await client.post(f"/api/public/trial-setup/{token}", json=SETUP)
    assert res.status_code == 200, res.text
    assert res.json()["data"]["access_token"]

    org_id = req.organization_id
    db_session.expire_all()
    org = await db_session.get(Organization, org_id)
    assert org.name == "Initech" and org.brand_color == "#FF5500" and org.email_sender_name == "Initech Talent"
    login = await client.post("/v1/auth/login", json={"email": "bill@initech-example.com", "password": "Initech2026"})
    assert login.status_code == 200
    # The link works once.
    assert (await client.get(f"/api/public/trial-setup/{token}")).status_code == 404


async def test_setup_link_expires_with_the_trial(client, db_session, plans, super_admin_headers):
    req, _ = await _approved(client, db_session, super_admin_headers)
    await _set_trial_end(db_session, req, datetime.now(timezone.utc) - timedelta(minutes=1))
    res = await client.get(f"/api/public/trial-setup/{req.setup_token}")
    assert res.status_code == 410 and "info@hybent.com" in res.json()["message"]


def test_org_emails_carry_the_org_name_and_reply_to(monkeypatch):
    """send_email uses the signed-in organization's branding for the From name and Reply-To."""
    import app.services.email_service as es

    calls = []
    monkeypatch.setattr(es.settings, "resend_api_key", "re_test")
    monkeypatch.setattr(es, "_send_resend", lambda to, subject, html, from_name=None, reply_to=None: calls.append((from_name, reply_to)))
    token = es.email_branding.set(es.EmailBranding(sender_name="Initech Talent", reply_to="bill@initech-example.com"))
    try:
        assert REAL_SEND_EMAIL("c@x-example.com", "Interview", "<p>hi</p>")
    finally:
        es.email_branding.reset(token)
    assert calls == [("Initech Talent", "bill@initech-example.com")]
    assert es._from_header("Initech Talent").startswith("Initech Talent <")


async def test_short_form_without_company_details_is_accepted(client, db_session):
    short = {k: v for k, v in VALID.items() if k not in ("company_size", "industry", "country", "hires_per_year")}
    token = await _verify(client, short["contact_email"])
    res = await client.post("/api/public/trial-request", json={**short, "email_verification_token": token})
    assert res.status_code == 200, res.text
    req = (await db_session.execute(select(TrialRequest))).scalar_one()
    assert req.company_size is None and req.industry is None and req.country is None


async def test_testing_mode_allows_repeat_requests(client, db_session, monkeypatch):
    monkeypatch.setattr(trial_service.settings, "trial_testing_mode", True)
    assert (await _submit(client)).status_code == 200
    first = (await db_session.execute(select(TrialRequest))).scalar_one()
    res = await _submit(client, contact_email="peter@initech-example.com")
    assert res.status_code == 200, res.text
    await db_session.refresh(first)
    assert first.status == "rejected" and "[testing]" in first.admin_notes


async def test_copilot_needs_force_on_standard(client, db_session, plans, organization, super_admin_headers):
    db_session.add(CompanySubscription(
        organization_id=organization.id, plan_id=plans["Standard"].id, status="active",
        current_period_end=datetime.now(timezone.utc) + timedelta(days=10),
    ))
    await db_session.commit()
    url = f"/v1/super-admin/clients/{organization.id}/flags"
    info = (await client.get(f"/v1/super-admin/clients/{organization.id}/plan-features", headers=super_admin_headers)).json()["data"]
    assert info == {"plan_name": "Standard", "copilot_in_plan": False}

    res = await client.put(url, headers=super_admin_headers, json={"copilot": True, "ai": True})
    assert res.status_code == 409 and res.json()["details"]["code"] == "copilot_requires_pro"
    assert (await client.put(url + "?force=true", headers=super_admin_headers, json={"copilot": True})).status_code == 200
    # Already on: saving other switches doesn't ask again; switching off is always allowed.
    assert (await client.put(url, headers=super_admin_headers, json={"copilot": True, "ai": False})).status_code == 200
    assert (await client.put(url, headers=super_admin_headers, json={"copilot": False})).status_code == 200
    audit = (await db_session.execute(select(SuperAdminAuditLog).where(SuperAdminAuditLog.action.like("Force-enabled%")))).scalars().all()
    assert len(audit) == 1


async def test_pro_client_switches_copilot_freely(client, db_session, plans, organization, super_admin_headers):
    db_session.add(CompanySubscription(
        organization_id=organization.id, plan_id=plans["Pro"].id, status="active",
        current_period_end=datetime.now(timezone.utc) + timedelta(days=10),
    ))
    await db_session.commit()
    url = f"/v1/super-admin/clients/{organization.id}/flags"
    assert (await client.put(url, headers=super_admin_headers, json={"copilot": False})).status_code == 200
    assert (await client.put(url, headers=super_admin_headers, json={"copilot": True})).status_code == 200
