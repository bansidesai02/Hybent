"""
Resume parsing service using Groq API.
Extracts structured data from PDF/DOCX resumes.
"""
import io
import json
import logging
import re
from datetime import datetime
from typing import Optional

import pdfplumber
import docx
import subprocess
import time
import uuid
from app.services.groq_client import SafeGroq as Groq, get_best_groq_model
from pydantic import BaseModel, Field, model_validator
from fastapi import BackgroundTasks, HTTPException

from app.core.config import settings
from app.services.ai_usage_tracker import log_ai_usage
from app.services.ai_metering import ai_feature

logger = logging.getLogger(__name__)

groq_client = Groq(api_key=settings.groq_api_key) if settings.groq_api_key else None

try:
    import google.generativeai as genai
    if settings.gemini_api_key:
        genai.configure(api_key=settings.gemini_api_key)
except Exception as exc:
    logger.warning(f"Gemini configuration error: {exc}")
    genai = None

GEMINI_FALLBACK_MODEL = "gemini-2.5-flash"


class NotAResumeError(HTTPException):
    """Raised when a document fails resume validation — unreadable, or
    positively classified as not a resume/CV (invoice, ID, certificate,
    etc.). Subclasses HTTPException so the manual-upload routers (which let
    parse_resume's exceptions propagate straight to a 400 response) keep
    working unchanged. Callers that need to tell "this genuinely isn't a
    resume" apart from a transient parsing failure — e.g. email ingestion,
    which should skip the former but still fall back gracefully on the
    latter — catch this type specifically."""

    def __init__(self, detail: str = "Invalid document. Please upload a valid professional resume/CV."):
        super().__init__(status_code=400, detail=detail)


class ProtectedDocumentError(NotAResumeError):
    """The PDF is locked with a password needed just to open it. Resumes are
    practically never sent like that; bank, broker and exchange statements
    (locked with a PAN or date of birth) routinely are — so email ingestion
    skips these rather than keeping them for review."""

    def __init__(self):
        super().__init__("This PDF is password-protected — remove the password and upload it again.")


class UnreadableDocumentError(NotAResumeError):
    """No text could be extracted, even with OCR — a very poor scan, a photo,
    or a password-protected file. Distinct from "read it, and it isn't a
    resume": nothing is known about the content, so email ingestion keeps the
    file for a human to review instead of discarding it. Still a
    NotAResumeError (400) for manual upload, with a message that says why."""

    def __init__(self):
        super().__init__(
            "We couldn't read any text from this file — it may be a low-quality scan or "
            "password-protected. Try a text-based PDF or DOCX."
        )


# ─── Pydantic schema ────────────────────────────────────────────────────────────

def _drop_nulls(data):
    """The LLM writes `null` for anything a résumé leaves out — a role with
    no dates, a missing name. Let those fall back to the field's default
    instead of failing validation: one null used to reject the whole parse,
    and the résumé then silently degraded to the regex fallback (no
    experience, no education, the name read as "Resume")."""
    if isinstance(data, dict):
        return {k: v for k, v in data.items() if v is not None}
    return data


class ExperienceEntry(BaseModel):
    title: str = ""
    company: str = ""
    duration: str = ""
    description: str = ""

    @model_validator(mode="before")
    @classmethod
    def _tolerate_nulls(cls, data):
        return _drop_nulls(data)


class ParsedResume(BaseModel):
    full_name: str = "Unknown"
    email: Optional[str] = None
    phone: Optional[str] = None
    location: Optional[str] = None
    current_title: Optional[str] = None
    current_company: Optional[str] = None
    years_experience: Optional[float] = None
    experience_years: Optional[str] = None
    summary: Optional[str] = None
    linkedin_url: Optional[str] = None
    github_url: Optional[str] = None
    portfolio_url: Optional[str] = None
    skills: list[str] = Field(default_factory=list)
    education: list[dict] = Field(default_factory=list)
    experience: list[ExperienceEntry] = Field(default_factory=list)
    projects: list[dict] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)
    languages: list[str] = Field(default_factory=list)

    @model_validator(mode="before")
    @classmethod
    def _tolerate_nulls(cls, data):
        data = _drop_nulls(data)
        if isinstance(data, dict):
            # Nulls inside lists too, e.g. "skills": ["Python", null].
            for key in ("skills", "certifications", "languages", "education", "experience", "projects"):
                if isinstance(data.get(key), list):
                    data[key] = [item for item in data[key] if item is not None]
        return data


# ─── Prompt ─────────────────────────────────────────────────────────────────────

PARSE_PROMPT = """
You are an expert resume parser and technical recruiter AI. Extract structured information from the resume text below.

For the 'summary', do NOT copy the candidate's objective. Analyze their entire resume (skills, experience, and projects) and write a comprehensive, original 3-4 sentence summary of their profile, technical strengths, and potential fit.

For 'experience', extract EVERY job entry with its exact duration string as written (e.g. "Jan 2020 - Dec 2022", "2019 - Present", "3 years 2 months").

Return ONLY a valid JSON object with this exact structure:
{
  "full_name": "string",
  "email": "string or null",
  "phone": "string or null",
  "location": "string or null",
  "current_title": "string or null",
  "current_company": "string or null",
  "years_experience": 0.0,
  "experience_years": "string (e.g. '3.5 Years') or null",
  "summary": "AI-generated analysis of the full resume",
  "linkedin_url": "string or null",
  "github_url": "string or null",
  "portfolio_url": "string or null",
  "skills": ["skill1", "skill2", ...],
  "education": [
    {"degree": "string", "institution": "string", "year": number or null}
  ],
  "experience": [
    {
      "title": "string",
      "company": "string",
      "duration": "exact duration string from resume e.g. Jan 2020 - Dec 2022",
      "description": "string"
    }
  ],
  "projects": [
    {
      "name": "string",
      "description": "string",
      "technologies": ["tech1", "tech2"]
    }
  ],
  "certifications": ["cert1", ...],
  "languages": ["English", ...]
}

Resume text:
"""

JD_PARSE_PROMPT = """
You are an expert technical recruiter and AI assistant. Extract structured job details from the Job Description text below.

CRITICAL for required_skills: Extract ONLY actual skill keywords, tool names, technologies, domain competencies, and certifications.
DO NOT include sentences, requirements text, or descriptions like "Bachelor's degree in..." or "Strong communication...".
Good examples: ["Python", "Django", "FastAPI", "SQL", "AWS", "Sales", "CRM", "Lead Generation", "B2B Sales", "LinkedIn Outreach", "Excel"]
Bad examples: ["Bachelor's degree in Business", "Strong communication skills", "Knowledge of sales processes"]

Return ONLY a valid JSON object with this exact structure (use null if not found):
{
  "title": "string (the job title)",
  "location": "string or null",
  "min_experience_years": number or null,
  "key_responsibilities": ["resp1", "resp2", ...],
  "required_skills": ["skill_keyword1", "skill_keyword2", ...],
  "description": "A clean, concise markdown version of the main job description, responsibilities, and qualifications."
}
"""


# ─── Text extraction ─────────────────────────────────────────────────────────────

# Below this many characters a document has no usable text layer — it is a
# scan, a photo, or protected — and the next extractor (finally OCR) is tried.
MIN_TEXT_CHARS = 150

# OCR is the slow path (~1-3s a page), so it stops after the first pages; a
# resume's identity, experience and skills are on page one or two anyway.
OCR_MAX_PAGES = 5
OCR_DPI = 200


def _pdf_text_pdfplumber(content: bytes) -> str:
    """Layout-preserving; handles tables and multi-column resumes best."""
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        parts = [page.extract_text(x_tolerance=3, y_tolerance=3) or "" for page in pdf.pages]
    return "\n".join(p for p in parts if p).strip()


def _pdf_text_pdfium(content: bytes) -> str:
    """Chrome's PDF engine (a pdfplumber dependency) — opens many files
    pdfminer chokes on (odd exporters, broken xref tables)."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(content)
    try:
        parts = []
        for page in pdf:
            textpage = page.get_textpage()
            parts.append(textpage.get_text_range())
            textpage.close()
            page.close()
        return "\n".join(parts).strip()
    finally:
        pdf.close()


def _pdf_text_pypdf2(content: bytes) -> str:
    """Also opens PDFs that are encrypted with only an owner password (no
    password needed to read) — common for resumes exported "protected"."""
    from PyPDF2 import PdfReader

    reader = PdfReader(io.BytesIO(content))
    if reader.is_encrypted:
        reader.decrypt("")
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


_ocr_unavailable_logged = False


def _pdf_text_ocr(content: bytes) -> str:
    """Tesseract OCR over page images, for scanned or image-only resumes.
    Pages are rendered with pypdfium2 — no extra system packages beyond
    tesseract itself (installed in the Dockerfile)."""
    global _ocr_unavailable_logged
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
    except Exception as e:
        if not _ocr_unavailable_logged:
            logger.warning(f"OCR unavailable (tesseract not installed?): {type(e).__name__}: {e}")
            _ocr_unavailable_logged = True
        return ""

    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(content)
    try:
        parts = []
        for index in range(min(len(pdf), OCR_MAX_PAGES)):
            page = pdf[index]
            # Grayscale at 200 DPI: accurate enough for print, ~4MB a page.
            image = page.render(scale=OCR_DPI / 72, grayscale=True).to_pil()
            parts.append(pytesseract.image_to_string(image, lang="eng"))
            page.close()
        return "\n".join(parts).strip()
    finally:
        pdf.close()


_PDF_EXTRACTORS = (
    ("pdfplumber", _pdf_text_pdfplumber),
    ("pdfium", _pdf_text_pdfium),
    ("pypdf2", _pdf_text_pypdf2),
    ("ocr", _pdf_text_ocr),
)


def is_password_protected_pdf(content: bytes) -> bool:
    """True when the PDF needs a password just to be opened. Owner-password
    PDFs ("protected" against editing, readable by anyone) return False."""
    try:
        from PyPDF2 import PdfReader

        reader = PdfReader(io.BytesIO(content))
        return bool(reader.is_encrypted) and not reader.decrypt("")
    except Exception:
        return False


def extract_text_from_pdf(content: bytes) -> str:
    """Extract a resume PDF's text, trying each extractor in turn until one
    yields a usable amount: three text-layer readers, then OCR for scans.

    Every failure is logged with its exception type — pdfminer's errors often
    have an empty message, which is why failures used to log as just
    "pdfplumber extraction failed: " with no clue what went wrong."""
    best = ""
    for name, extract in _PDF_EXTRACTORS:
        try:
            text = extract(content)
        except Exception as e:
            logger.warning(f"PDF text extraction via {name} failed: {type(e).__name__}: {str(e) or 'no message'}")
            continue
        if len(text) >= MIN_TEXT_CHARS:
            if name != "pdfplumber":
                logger.info(f"PDF text extracted via {name} ({len(text)} chars)")
            return text
        if len(text) > len(best):
            best = text
    logger.warning(f"No extractor found a usable text layer in this PDF (best: {len(best)} chars)")
    return best


def extract_text_from_docx(content: bytes) -> str:
    """Extract plain text from DOCX bytes including tables and headers/footers."""
    try:
        doc = docx.Document(io.BytesIO(content))
        full_text = []

        # 1. Headers/Footers
        for section in doc.sections:
            if section.header:
                for p in section.header.paragraphs:
                    if p.text.strip():
                        full_text.append(p.text)
            if section.footer:
                for p in section.footer.paragraphs:
                    if p.text.strip():
                        full_text.append(p.text)

        # 2. Main Paragraphs
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text)

        # 3. Tables (Crucial for many resume layouts)
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    full_text.append(" | ".join(row_text))

        return "\n".join(full_text).strip()
    except Exception as e:
        logger.error(f"DOCX extraction failed: {e}")
        return ""


def extract_text_from_doc(content: bytes) -> str:
    """Extract plain text from legacy binary .doc bytes using antiword."""
    try:
        # Create a temporary file to hold the content for antiword
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".doc", delete=True) as tmp:
            tmp.write(content)
            tmp.flush()
            
            # Run antiword through subprocess
            # -w 0 means no line wrapping
            result = subprocess.run(
                ["antiword", "-w", "0", tmp.name],
                capture_output=True,
                text=True,
                check=True
            )
            return result.stdout.strip()
    except subprocess.CalledProcessError as e:
        logger.error(f"antiword extraction failed: {e}")
        return ""
    except FileNotFoundError:
        logger.error("antiword not found on system. Legacy .doc support is disabled.")
        return ""
    except Exception as e:
        logger.error(f"Legacy .doc extraction failed: {e}")
        return ""


def _detect_file_type(file_content: bytes, content_type: str, filename: str = "") -> str:
    """Returns 'pdf', 'docx', 'doc', or 'unknown'."""
    ct = content_type.lower()
    fn = filename.lower()
    
    # Check by extension first for clarity
    if fn.endswith(".pdf"):
        return "pdf"
    if fn.endswith(".docx"):
        return "docx"
    if fn.endswith(".doc"):
        return "doc"

    # Check by magic bytes / content-type
    if "pdf" in ct or file_content[:4] == b"%PDF":
        return "pdf"
    
    if "docx" in ct or "openxmlformats" in ct:
        return "docx"
    
    if "msword" in ct or "document" in ct:
        # Many sources report .docx as generic "document"
        if file_content[:2] == b"PK": # Zip signature (DOCX is a zip)
            return "docx"
        return "doc"
    
    if "zip" in ct or file_content[:2] == b"PK":
        return "docx" # Assume it's a DOCX/OPC file if it's a zip sent to a resume parser
    
    # Signature for legacy OLE binary doc files
    if file_content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "doc"

    return "unknown"


# ─── Years experience calculation ────────────────────────────────────────────────

# One date inside a résumé duration string. Résumés write these many ways:
# "Jan 2020", "January, 2020", "Apr - 2021", "Sept2020", "06/2024", "2019".
_MONTHS = {"jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
           "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12}
_DATE_TOKEN = re.compile(
    r"\b(?:"
    r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?[\s,.'\u2019/\-]*((?:19|20)\d{2})"  # month name
    r"|(0?[1-9]|1[0-2])\s*[/.\-]\s*((?:19|20)\d{2})"  # numeric MM/YYYY
    r"|((?:19|20)\d{2})"  # bare year
    r")\b",
    re.IGNORECASE,
)
_ONGOING = re.compile(r"\b(present|current|currently|now|today|till date|to date|ongoing)\b", re.IGNORECASE)


def _months_in_date_range(duration: str) -> Optional[int]:
    """Inclusive month count of a "<start> - <end>" duration, or None when
    fewer than two dates can be read. An end of "Present"/"Current" is this
    month. A bare year has no month: as a start it counts from January, as
    an end through December. Mirrors formatExperienceDuration on the
    frontend, which renders the per-role badge from the same strings."""
    now = datetime.now()
    norm = _ONGOING.sub(f"{now.strftime('%b')} {now.year}", duration)

    dates: list[tuple[int, Optional[int]]] = []
    for m in _DATE_TOKEN.finditer(norm):
        if m.group(1):
            dates.append((int(m.group(2)), _MONTHS[m.group(1)[:3].lower()]))
        elif m.group(3):
            dates.append((int(m.group(4)), int(m.group(3))))
        else:
            dates.append((int(m.group(5)), None))
    if len(dates) < 2:
        return None

    (y1, m1), (y2, m2) = dates[0], dates[-1]
    months = (y2 - y1) * 12 + ((m2 or 12) - (m1 or 1)) + 1
    return months if months > 0 else None


def calculate_years_from_experience(experience_list: list) -> tuple[Optional[float], Optional[str]]:
    """
    Calculate total years of experience from experience entries.
    Never trust the LLM's years_experience — compute it accurately from dates.
    Handles short internships (e.g. 1 month / exact dates) correctly.
    Returns: (years_float, experience_years_str)
    """
    total_months = 0

    for exp in experience_list:
        if hasattr(exp, 'model_dump'):
            exp = exp.model_dump()
        duration = str(exp.get('duration', '') or '').strip()
        if not duration:
            continue

        # Pattern: "X years Y months" or "X months" explicitly stated (supporting decimals like 2.5)
        years_match = re.search(r'(\d+(?:\.\d+)?)\s*(?:yr|year|years?)', duration, re.IGNORECASE)
        months_match = re.search(r'(\d+(?:\.\d+)?)\s*(?:mo|month|months?)', duration, re.IGNORECASE)
        if years_match or months_match:
            if years_match:
                total_months += float(years_match.group(1)) * 12
            if months_match:
                total_months += float(months_match.group(1))
            continue

        # Pattern: date range. Needs at least two resolvable dates.
        months = _months_in_date_range(duration)
        if months:
            total_months += months
        # A single bare year ("2024") with no range and no explicit "X years/
        # months" text isn't a duration — it's a date with the length undetermined.
        # Guessing "~1 month" here is what produced misleading 0.1-year entries;
        # an entry we can't confidently resolve contributes nothing, rather than
        # a fabricated number.

    if total_months == 0:
        return None, None

    years_float = round(total_months / 12.0, 1)
    
    if total_months < 12:
        display_months = int(total_months) if float(total_months).is_integer() else total_months
        experience_years_str = f"{display_months} {'Month' if display_months == 1 else 'Months'}"
    else:
        display_years = int(years_float) if float(years_float).is_integer() else years_float
        experience_years_str = f"{display_years} {'Year' if display_years == 1 else 'Years'}"
    
    return years_float, experience_years_str


def apply_experience_fields(candidate, parsed: dict) -> None:
    """
    Set `years_experience` and `experience_years` on a Candidate together, from
    a `parse_resume()` result dict.

    `parsed["years_experience"]`/`parsed["experience_years"]` are already
    computed as a matched pair by `parse_resume()` (via
    `calculate_years_from_experience`) — the bug this exists to prevent is an
    endpoint copying one of the two onto the model and forgetting the other,
    which is what previously left some candidates with a bare, unlabeled
    `years_experience` float and no `experience_years` string to disambiguate
    it. Every candidate-creating/updating endpoint should call this instead of
    assigning the two fields itself.
    """
    if parsed.get("years_experience") is not None:
        candidate.years_experience = parsed["years_experience"]
        candidate.experience_years = parsed.get("experience_years")


# ─── Fallback ────────────────────────────────────────────────────────────────────

def _regex_fallback(text: str) -> dict:
    parsed = {
        "raw_text": text[:500],
        "full_name": None,
        "email": None,
        "phone": None,
        "location": None,
        "current_title": None,
        "current_company": None,
        "years_experience": None,
        "summary": None,
        "skills": [],
    }

    email_match = re.search(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', text, re.IGNORECASE)
    if email_match:
        parsed["email"] = email_match.group(0).lower().strip()

    phone_match = re.search(r'(\+\d{1,2}\s)?\(?\d{3}\)?[\s.-]\d{3}[\s.-]\d{4}', text)
    if phone_match:
        parsed["phone"] = phone_match.group(0)

    lines = [line.strip() for line in text.split('\n') if line.strip()]
    if lines and len(lines[0]) < 50:
        parsed["full_name"] = lines[0].title()

    # Extract common technical skill keywords from raw text
    common_skills = [
        "Python", "JavaScript", "TypeScript", "React", "Node.js", "Java", "C++", "C#",
        "SQL", "PostgreSQL", "MySQL", "MongoDB", "Redis", "AWS", "Docker", "Kubernetes",
        "HTML", "CSS", "Git", "FastAPI", "Django", "Flask", "Express", "Next.js",
        "Tailwind", "REST API", "GraphQL", "Go", "Rust", "PHP", "Ruby", "Angular", "Vue"
    ]
    text_lower = text.lower()
    extracted_skills = []
    for skill in common_skills:
        pattern = r'\b' + re.escape(skill.lower()) + r'\b'
        if re.search(pattern, text_lower):
            extracted_skills.append(skill)
    parsed["skills"] = extracted_skills

    return parsed


# ─── Groq call with Pydantic validation + retry ──────────────────────────────────

def _call_groq_with_retry(
    text: str,
    background_tasks: Optional[BackgroundTasks] = None,
    user_id: Optional[uuid.UUID] = None,
    organization_id: Optional[uuid.UUID] = None
) -> Optional[dict]:
    """
    Call Groq llama-3.3-70b-versatile, validate with Pydantic, retry once on failure.
    Returns validated dict with years_experience calculated from dates.
    """
    last_error = None
    start_time = time.time()
    p_tokens, c_tokens, t_tokens = 0, 0, 0
    status = "success"

    for attempt in range(2):
        try:
            messages = [
                {"role": "system", "content": "You are a helpful assistant that outputs ONLY valid JSON."},
                {"role": "user", "content": PARSE_PROMPT + text[:8000]},
            ]
            if attempt == 1 and last_error:
                messages[1]["content"] += f"\n\nIMPORTANT: Previous attempt failed validation: {last_error}. Fix the JSON structure."

            completion = groq_client.chat.completions.create(
                messages=messages,
                model=get_best_groq_model(groq_client),
                response_format={"type": "json_object"},
                temperature=0.1,
            )
            
            if hasattr(completion, 'usage'):
                p_tokens = completion.usage.prompt_tokens
                c_tokens = completion.usage.completion_tokens
                t_tokens = completion.usage.total_tokens

            content = completion.choices[0].message.content
            if not content:
                last_error = "Empty response"
                continue

            raw = json.loads(content)
            validated = ParsedResume(**raw)
            result = validated.model_dump()

            # Always compute from dates because LLMs are not reliable at date math.
            computed_years, computed_str = calculate_years_from_experience(validated.experience)
            if computed_years is not None:
                result["years_experience"] = computed_years
                result["experience_years"] = computed_str
            else:
                # Fallback to LLM values
                if result.get("years_experience") is not None and not result.get("experience_years"):
                    result["experience_years"] = f"{result['years_experience']} Years"

            # Serialize experience to plain dicts
            result["experience"] = [
                e.model_dump() if hasattr(e, 'model_dump') else e
                for e in validated.experience
            ]
            return result

        except Exception as e:
            status = "failure"
            last_error = str(e)
            if attempt == 0:
                logger.warning(f"Groq parse attempt 1 failed ({e}), retrying...")
            else:
                logger.error(f"Groq parse attempt 2 failed ({e})")
        finally:
            if attempt == 1 or status == "success":
                duration_ms = (time.time() - start_time) * 1000
                if background_tasks:
                    background_tasks.add_task(
                        log_ai_usage,
                        provider="Groq",
                        model="llama-3.3-70b-versatile",
                        feature="resume_parsing",
                        prompt_tokens=p_tokens,
                        completion_tokens=c_tokens,
                        total_tokens=t_tokens,
                        duration_ms=duration_ms,
                        status=status,
                        error_detail=last_error if status == "failure" else None,
                        user_id=user_id,
                        organization_id=organization_id
                    )

    return None


def _verify_is_resume_with_keywords(text: str) -> bool:
    text_lower = text.lower()

    # Exclude government IDs and unrelated documents by checking for specific words
    blacklist_patterns = [
        r"government of india", r"permanent account number", r"aadhaar", r"income tax department",
        r"republic of india", r"\bpassport\b", r"\bdriving license\b", r"\bbill of entry\b",
        r"\bmarksheet\b", r"certificate of completion", r"academic transcript",
        r"\bpurchase order\b", r"\bpan card\b", r"unique identification authority", r"tax invoice"
    ]
    for pattern in blacklist_patterns:
        if re.search(pattern, text_lower):
            logger.info(f"Document validation failed: matched blacklist pattern '{pattern}'")
            return False

    # Many resume templates (esp. Canva-style designs exported to PDF) render
    # section headings with wide letter-spacing, e.g. "E D U C A T I O N".
    # pdfplumber preserves that as literal spaces between characters, so a
    # plain `"education" in text` substring check never matches the heading
    # even though the document plainly is a resume — this legitimate resume
    # was being rejected for exactly that reason. Matching against a
    # whitespace-collapsed copy of the text catches both the normal case and
    # this letter-spaced-heading case, without weakening the check (a random
    # non-resume document collapsing to contain these specific keywords by
    # coincidence is effectively impossible).
    text_collapsed = re.sub(r"\s+", "", text_lower)

    # Resumes typically contain a combination of keywords from different sections
    resume_keywords = [
        "experience", "education", "skills", "projects", "employment",
        "summary", "work history", "academic", "qualifications", "curriculum vitae", "resume"
    ]
    matches = sum(
        1 for kw in resume_keywords
        if kw in text_lower or kw.replace(" ", "") in text_collapsed
    )
    # If it has at least 2 common resume section keywords, consider it a resume
    is_res = matches >= 2
    logger.info(f"Document keyword validation result: matches={matches}, is_resume={is_res}")
    return is_res


_DOCUMENT_CLASSIFIER_PROMPT = """
You are an expert AI document classifier. Analyze the text below and determine if it is a genuine professional resume or curriculum vitae (CV).

A genuine resume/CV MUST contain details about a person's professional history, such as their work experience, professional skills, education, or project history.
Section headings may have unusual letter-spacing (e.g. "E D U C A T I O N") due to PDF export artifacts — do not treat that as a reason to reject.

You MUST reject documents that are NOT resumes, including:
- Identity cards / government documents (e.g. Aadhaar, PAN, Passports, Driving Licenses, SSNs)
- Certificates (e.g. course completion, degree certificates)
- Academic transcripts / Marksheets
- Business documents (e.g. Invoices, receipts, purchase orders, offer letters, employment contracts)
- Random letters, articles, essays, or unrelated text.

Return ONLY a valid JSON object with the following structure:
{
  "is_resume": true or false,
  "reason": "a brief explanation of your decision"
}

Text to analyze:
"""


def _verify_is_resume_with_gemini(text: str) -> Optional[bool]:
    """Secondary classifier — only reached when Groq (all configured keys,
    every model) has failed. Returns None (not False!) on any failure so the
    caller falls through to the keyword check instead of wrongly treating
    'Gemini unavailable' as 'not a resume'."""
    if genai is None or not settings.gemini_api_key:
        return None
    try:
        model = genai.GenerativeModel(GEMINI_FALLBACK_MODEL)
        response = model.generate_content(
            _DOCUMENT_CLASSIFIER_PROMPT + text[:4000],
            generation_config={"response_mime_type": "application/json", "temperature": 0.1},
        )
        if response and response.text:
            cleaned = response.text.strip()
            if cleaned.startswith("```"):
                cleaned = cleaned.strip("`").removeprefix("json").strip()
            res = json.loads(cleaned)
            is_res = res.get("is_resume", False)
            logger.info(f"Gemini document validation result: is_resume={is_res}, reason={res.get('reason')}")
            return is_res
    except Exception as e:
        logger.warning(f"Gemini document validation failover error: {e}")
    return None


def _verify_is_resume_with_llm(text: str) -> bool:
    if not groq_client:
        gemini_result = _verify_is_resume_with_gemini(text)
        if gemini_result is not None:
            return gemini_result
        return _verify_is_resume_with_keywords(text)

    try:
        completion = groq_client.chat.completions.create(
            messages=[
                {"role": "system", "content": "You are a helpful assistant that outputs ONLY valid JSON."},
                {"role": "user", "content": _DOCUMENT_CLASSIFIER_PROMPT + text[:4000]},
            ],
            model=get_best_groq_model(groq_client),
            response_format={"type": "json_object"},
            temperature=0.1,
            max_tokens=100,
        )
        content = completion.choices[0].message.content
        if content:
            res = json.loads(content)
            is_res = res.get("is_resume", False)
            logger.info(f"Document validation result: is_resume={is_res}, reason={res.get('reason')}")
            return is_res
    except Exception as e:
        logger.error(f"Error during document validation (Groq): {e}. Trying Gemini failover...")

    gemini_result = _verify_is_resume_with_gemini(text)
    if gemini_result is not None:
        return gemini_result

    return _verify_is_resume_with_keywords(text)


# ─── Public API ──────────────────────────────────────────────────────────────────

@ai_feature("resume_parsing")
async def parse_resume(
    file_content: bytes, 
    content_type: str, 
    filename: str = "",
    background_tasks: Optional[BackgroundTasks] = None,
    user_id: Optional[uuid.UUID] = None,
    organization_id: Optional[uuid.UUID] = None
) -> dict:
    """Main entry point: parse resume bytes and return structured dict."""
    file_type = _detect_file_type(file_content, content_type, filename)
    
    if file_type == "pdf":
        text = extract_text_from_pdf(file_content)
    elif file_type == "docx":
        text = extract_text_from_docx(file_content)
    elif file_type == "doc":
        text = extract_text_from_doc(file_content)
    else:
        logger.warning(f"Unsupported resume format or content type: {content_type}")
        raise NotAResumeError()

    if not text or len(text.strip()) < MIN_TEXT_CHARS:
        if file_type == "pdf" and is_password_protected_pdf(file_content):
            logger.info(f"Resume {filename!r} is a password-protected PDF")
            raise ProtectedDocumentError()
        logger.warning(f"Could not extract usable text from resume {filename!r} ({len((text or '').strip())} chars)")
        raise UnreadableDocumentError()

    # Perform strict document validation
    if groq_client:
        is_resume = _verify_is_resume_with_llm(text)
    else:
        is_resume = _verify_is_resume_with_keywords(text)

    if not is_resume:
        raise NotAResumeError()

    if organization_id:
        from app.services.ai_credit_service import AICreditsService
        await AICreditsService.check_credits_available(None, organization_id, "resume_parsing")

    if not groq_client:
        logger.warning("No Groq API key configured — using regex fallback")
        return _regex_fallback(text)

    try:
        result = _call_groq_with_retry(
            text,
            background_tasks=background_tasks,
            user_id=user_id,
            organization_id=organization_id
        )
        if result:
            return result
        return _regex_fallback(text)
    except Exception as e:
        error_msg = str(e).lower()
        if "rate balance" in error_msg or "429" in error_msg:
            logger.error(f"Groq API Quota Exceeded: {e}")
            raise HTTPException(
                status_code=429,
                detail="Groq AI service is currently rate limited. Please try again soon."
            )
        logger.error(f"Groq parsing failed: {e}")
        return _regex_fallback(text)


@ai_feature("jd_parsing")
async def parse_jd(
    file_bytes: bytes, 
    content_type: str,
    background_tasks: Optional[BackgroundTasks] = None,
    user_id: Optional[uuid.UUID] = None,
    organization_id: Optional[uuid.UUID] = None
) -> dict:
    """Extracts job details using Groq."""
    if "pdf" in content_type:
        text = extract_text_from_pdf(file_bytes)
    elif "docx" in content_type or "document" in content_type:
        text = extract_text_from_docx(file_bytes)
    elif "text/plain" in content_type or "text" in content_type:
        text = file_bytes.decode('utf-8', errors='ignore')
    else:
        text = file_bytes.decode('utf-8', errors='ignore')

    if not text.strip():
        raise ValueError("No text could be extracted from the file.")

    if not groq_client:
        logger.warning("Groq API key not configured. Returning empty JD data.")
        return {
            "title": "Software Engineer",
            "location": None,
            "min_experience_years": None,
            "key_responsibilities": [],
            "required_skills": [],
            "description": text[:800],
        }

    start_time = time.time()
    p_tokens, c_tokens, t_tokens = 0, 0, 0
    status = "success"
    error_msg = None

    try:
        response = groq_client.chat.completions.create(
            messages=[
                {"role": "system", "content": "You are a helpful assistant that outputs ONLY valid JSON. " + JD_PARSE_PROMPT},
                {"role": "user", "content": text[:15000]},
            ],
            model=get_best_groq_model(groq_client),
            response_format={"type": "json_object"},
            temperature=0.1,
        )
        
        if hasattr(response, 'usage'):
            p_tokens = response.usage.prompt_tokens
            c_tokens = response.usage.completion_tokens
            t_tokens = response.usage.total_tokens

        content = response.choices[0].message.content
        if not content:
            raise ValueError("Empty response from AI")
        return json.loads(content.strip())
    except Exception as e:
        status = "failure"
        error_msg = str(e)
        logger.error(f"Error parsing JD with Groq: {e}")
        return {
            "title": "Unknown Title",
            "required_skills": [],
            "description": text[:500],
        }
    finally:
        duration_ms = (time.time() - start_time) * 1000
        if background_tasks:
            background_tasks.add_task(
                log_ai_usage,
                provider="Groq",
                model="llama-3.3-70b-versatile",
                feature="jd_parsing",
                prompt_tokens=p_tokens,
                completion_tokens=c_tokens,
                total_tokens=t_tokens,
                duration_ms=duration_ms,
                status=status,
                error_detail=error_msg,
                user_id=user_id,
                organization_id=organization_id
            )

@ai_feature("match_summary")
async def generate_match_summary(
    candidate_data: dict, 
    job_title: str, 
    job_skills: list[str], 
    score: float, 
    match_threshold: float,
    decision_reasons: list[str],
    background_tasks: Optional[BackgroundTasks] = None,
    user_id: Optional[uuid.UUID] = None,
    organization_id: Optional[uuid.UUID] = None
) -> str:
    """
    Generates a 2-3 sentence paragraph explaining why a candidate received their match score,
    focusing on their fit against the specific job requirements.
    """
    if not groq_client:
        return "AI analysis unavailable. Please refer to the specific bullet points above."

    shortlisted = score >= match_threshold
    status_text = "Shortlisted" if shortlisted else "Rejected / Needs Review"

    prompt = f"""
    You are an expert technical recruiter AI. Write a concise, professional 2-3 sentence summary explaining exactly why this candidate was {status_text} for the {job_title} role.
    
    Context:
    - Candidate Score: {score}% (Threshold: {match_threshold}%)
    - Candidate Skills: {', '.join(candidate_data.get('skills', [])[:15])}
    - Candidate Experience: {candidate_data.get('years_experience')} years
    - Required Job Skills: {', '.join(job_skills[:15])}
    
    Decision Breakdown:
    {chr(10).join(decision_reasons)}
    
    Guidelines:
    - Do NOT use bullet points. Write a single short paragraph.
    - Be direct but professional. Speak about the candidate in the third person.
    - Specifically mention what they matched well on OR what key skills/experience they are missing that caused the low score.
    - Keep it strictly under 50 words.
    """

    start_time = time.time()
    p_tokens, c_tokens, t_tokens = 0, 0, 0
    status = "success"
    error_msg = None

    try:
        response = groq_client.chat.completions.create(
            messages=[
                {"role": "system", "content": "You are an expert technical recruiter. Output only the requested summary paragraph."},
                {"role": "user", "content": prompt},
            ],
            model=get_best_groq_model(groq_client),
            temperature=0.3,
            max_tokens=150,
        )
        
        if hasattr(response, 'usage'):
            p_tokens = response.usage.prompt_tokens
            c_tokens = response.usage.completion_tokens
            t_tokens = response.usage.total_tokens

        content = response.choices[0].message.content
        if not content:
            return "Could not generate match summary."
        return content.strip()
    except Exception as e:
        status = "failure"
        error_msg = str(e)
        logger.error(f"Error generating match summary with Groq: {e}")
        return "Could not generate match summary."
    finally:
        duration_ms = (time.time() - start_time) * 1000
        if background_tasks:
            background_tasks.add_task(
                log_ai_usage,
                provider="Groq",
                model="llama-3.3-70b-versatile",
                feature="match_summary",
                prompt_tokens=p_tokens,
                completion_tokens=c_tokens,
                total_tokens=t_tokens,
                duration_ms=duration_ms,
                status=status,
                error_detail=error_msg,
                user_id=user_id,
                organization_id=organization_id
            )
